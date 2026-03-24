"""Tests for cover_fetcher.csv_handler"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.styles import numbers
import pytest

from cover_fetcher.csv_handler import (
    EXCEL_IMAGE_COLUMN,
    IMAGE_COLUMN,
    build_image_filename,
    ensure_image_column,
    iter_pending_rows,
    reserve_unique_filename,
    sanitize_filename,
    update_row,
    update_rows,
)


# ---------------------------------------------------------------------------
# sanitize_filename
# ---------------------------------------------------------------------------


def test_sanitize_removes_unsafe():
    result = sanitize_filename('AC/DC: "Back in Black"?')
    assert "/" not in result
    assert ":" not in result
    assert '"' not in result
    assert "?" not in result


def test_sanitize_collapses_spaces():
    result = sanitize_filename("Hello   World")
    assert "  " not in result


# ---------------------------------------------------------------------------
# build_image_filename
# ---------------------------------------------------------------------------


def test_build_image_filename_basic():
    name = build_image_filename("The Beatles", "Abbey Road", "LP")
    assert name.endswith(".jpg")
    assert "Beatles" in name
    assert "Abbey" in name
    assert "LP" in name
    assert "2024" not in name


def test_build_image_filename_unsafe_chars():
    name = build_image_filename("AC/DC", "Back/In/Black", "LP")
    assert "/" not in name


def test_reserve_unique_filename_adds_suffix_when_name_exists(tmp_path):
    (tmp_path / "3D_Fantasia_7.jpg").write_bytes(b"existing")
    reserved: set[str] = set()

    result = reserve_unique_filename("3D_Fantasia_7.jpg", tmp_path, reserved)

    assert result == "3D_Fantasia_7_2.jpg"


def test_reserve_unique_filename_avoids_batch_collisions(tmp_path):
    reserved: set[str] = set()

    first = reserve_unique_filename("3D_Fantasia_7.jpg", tmp_path, reserved)
    second = reserve_unique_filename("3D_Fantasia_7.jpg", tmp_path, reserved)

    assert first == "3D_Fantasia_7.jpg"
    assert second == "3D_Fantasia_7_2.jpg"


# ---------------------------------------------------------------------------
# ensure_image_column
# ---------------------------------------------------------------------------


def _write_simple_csv(path: Path, rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def _write_simple_xlsx(path: Path, headers: list[str], rows: list[list[str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()


def test_ensure_image_column_adds_column(tmp_path):
    csv_path = tmp_path / "collection.csv"
    _write_simple_csv(csv_path, [{"artist": "X", "title": "Y"}])
    ensure_image_column(csv_path)
    with csv_path.open() as fh:
        reader = csv.DictReader(fh)
        assert IMAGE_COLUMN in (reader.fieldnames or [])


def test_ensure_image_column_idempotent(tmp_path):
    csv_path = tmp_path / "collection.csv"
    _write_simple_csv(csv_path, [{"artist": "X", "title": "Y", IMAGE_COLUMN: ""}])
    ensure_image_column(csv_path)
    ensure_image_column(csv_path)
    with csv_path.open() as fh:
        reader = csv.DictReader(fh)
        cols = reader.fieldnames or []
    assert cols.count(IMAGE_COLUMN) == 1


def test_ensure_image_column_requires_existing_image_column_in_xlsx(tmp_path):
    xlsx_path = tmp_path / "collection.xlsx"
    _write_simple_xlsx(xlsx_path, ["Band", "Title"], [["A", "Song"]])

    with pytest.raises(ValueError, match="must contain an existing 'Image' column"):
        ensure_image_column(xlsx_path)


def test_ensure_image_column_accepts_existing_image_column_in_xlsx(tmp_path):
    xlsx_path = tmp_path / "collection.xlsx"
    _write_simple_xlsx(xlsx_path, ["Band", "Title", EXCEL_IMAGE_COLUMN], [["A", "Song", ""]])

    ensure_image_column(xlsx_path)


# ---------------------------------------------------------------------------
# iter_pending_rows
# ---------------------------------------------------------------------------


def test_iter_pending_rows_all_pending(tmp_path):
    csv_path = tmp_path / "collection.csv"
    rows = [
        {"artist": "A", "title": "T1", IMAGE_COLUMN: ""},
        {"artist": "B", "title": "T2", IMAGE_COLUMN: ""},
    ]
    _write_simple_csv(csv_path, rows)
    pending = list(iter_pending_rows(csv_path, tmp_path))
    assert len(pending) == 2
    assert pending[0][0] == 0
    assert pending[1][0] == 1


def test_iter_pending_rows_skips_done(tmp_path):
    csv_path = tmp_path / "collection.csv"
    # Create a fake image file
    (tmp_path / "cover.jpg").write_bytes(b"fake")
    rows = [
        {"artist": "A", "title": "T1", IMAGE_COLUMN: "cover.jpg"},
        {"artist": "B", "title": "T2", IMAGE_COLUMN: ""},
    ]
    _write_simple_csv(csv_path, rows)
    pending = list(iter_pending_rows(csv_path, tmp_path))
    assert len(pending) == 1
    assert pending[0][0] == 1


def test_iter_pending_rows_reprocesses_deleted_image(tmp_path):
    csv_path = tmp_path / "collection.csv"
    rows = [
        {"artist": "A", "title": "T1", IMAGE_COLUMN: "missing.jpg"},
    ]
    _write_simple_csv(csv_path, rows)
    # file does NOT exist
    pending = list(iter_pending_rows(csv_path, tmp_path))
    assert len(pending) == 1


def test_iter_pending_rows_parses_semicolon_delimited_csv(tmp_path):
    csv_path = tmp_path / "collection_semicolon.csv"
    csv_path.write_text(
        "Band;Title;year;FORMAT;Label;image_file\n"
        "2 Lux 25;Go !!! / Give Yourself;91;\"12\\\"\";Extreme;\n",
        encoding="utf-8",
    )

    pending = list(iter_pending_rows(csv_path, tmp_path))
    assert len(pending) == 1
    _, row = pending[0]
    assert row["Band"] == "2 Lux 25"
    assert row["Title"] == "Go !!! / Give Yourself"
    assert row["year"] == "91"
    assert row["FORMAT"] == '12"'
    assert row["Label"] == "Extreme"


def test_iter_pending_rows_xlsx_uses_image_column(tmp_path):
    xlsx_path = tmp_path / "collection.xlsx"
    (tmp_path / "done.jpg").write_bytes(b"ok")
    _write_simple_xlsx(
        xlsx_path,
        ["Band", "Title", EXCEL_IMAGE_COLUMN],
        [["A", "Done", "done.jpg"], ["B", "Pending", ""]],
    )

    pending = list(iter_pending_rows(xlsx_path, tmp_path))
    assert len(pending) == 1
    assert pending[0][0] == 1
    assert pending[0][1]["Band"] == "B"


# ---------------------------------------------------------------------------
# update_row
# ---------------------------------------------------------------------------


def test_update_row_sets_value(tmp_path):
    csv_path = tmp_path / "collection.csv"
    _write_simple_csv(csv_path, [
        {"artist": "A", "title": "T", IMAGE_COLUMN: ""},
        {"artist": "B", "title": "U", IMAGE_COLUMN: ""},
    ])
    update_row(csv_path, 0, "a_t_lp_2024.jpg")
    with csv_path.open() as fh:
        reader = csv.DictReader(fh)
        rows = list(reader)
    assert rows[0][IMAGE_COLUMN] == "a_t_lp_2024.jpg"
    assert rows[1][IMAGE_COLUMN] == ""


def test_update_row_out_of_range(tmp_path):
    csv_path = tmp_path / "collection.csv"
    _write_simple_csv(csv_path, [{"artist": "A", "title": "T", IMAGE_COLUMN: ""}])
    with pytest.raises(IndexError):
        update_row(csv_path, 99, "x.jpg")


def test_update_row_sets_value_in_xlsx_image_column(tmp_path):
    xlsx_path = tmp_path / "collection.xlsx"
    _write_simple_xlsx(
        xlsx_path,
        ["Band", "Title", EXCEL_IMAGE_COLUMN],
        [["A", "T", ""], ["B", "U", ""]],
    )

    update_row(xlsx_path, 0, "a_t_lp.jpg")

    wb = load_workbook(xlsx_path, data_only=True)
    try:
        ws = wb.active
        assert ws.cell(row=2, column=3).value == "a_t_lp.jpg"
        assert ws.cell(row=3, column=3).value in (None, "")
    finally:
        wb.close()


def test_update_rows_sets_multiple_values_in_csv(tmp_path):
    csv_path = tmp_path / "collection.csv"
    _write_simple_csv(csv_path, [
        {"artist": "A", "title": "T", IMAGE_COLUMN: ""},
        {"artist": "B", "title": "U", IMAGE_COLUMN: ""},
        {"artist": "C", "title": "V", IMAGE_COLUMN: ""},
    ])

    update_rows(csv_path, {0: "a.jpg", 2: "c.jpg"})

    with csv_path.open() as fh:
        rows = list(csv.DictReader(fh))
    assert rows[0][IMAGE_COLUMN] == "a.jpg"
    assert rows[1][IMAGE_COLUMN] == ""
    assert rows[2][IMAGE_COLUMN] == "c.jpg"


def test_update_rows_sets_multiple_values_in_xlsx(tmp_path):
    xlsx_path = tmp_path / "collection.xlsx"
    _write_simple_xlsx(
        xlsx_path,
        ["Band", "Title", EXCEL_IMAGE_COLUMN],
        [["A", "T", ""], ["B", "U", ""], ["C", "V", ""]],
    )

    update_rows(xlsx_path, {0: "a.jpg", 2: "c.jpg"})

    wb = load_workbook(xlsx_path, data_only=True)
    try:
        ws = wb.active
        assert ws.cell(row=2, column=3).value == "a.jpg"
        assert ws.cell(row=3, column=3).value in (None, "")
        assert ws.cell(row=4, column=3).value == "c.jpg"
    finally:
        wb.close()


def test_update_row_preserves_xlsx_cell_number_format(tmp_path):
    xlsx_path = tmp_path / "collection.xlsx"
    _write_simple_xlsx(
        xlsx_path,
        ["Band", "Title", EXCEL_IMAGE_COLUMN],
        [["A", "T", ""]],
    )

    wb = load_workbook(xlsx_path)
    try:
        ws = wb.active
        ws.cell(row=2, column=3).number_format = numbers.FORMAT_TEXT
        wb.save(xlsx_path)
    finally:
        wb.close()

    update_row(xlsx_path, 0, "a_t_lp.jpg")

    wb = load_workbook(xlsx_path)
    try:
        ws = wb.active
        assert ws.cell(row=2, column=3).value == "a_t_lp.jpg"
        assert ws.cell(row=2, column=3).number_format == numbers.FORMAT_TEXT
    finally:
        wb.close()


def test_update_row_inherits_row_style_when_target_cell_unstyled(tmp_path):
    xlsx_path = tmp_path / "collection.xlsx"
    _write_simple_xlsx(
        xlsx_path,
        ["Band", "Title", EXCEL_IMAGE_COLUMN],
        [["A", "T", ""]],
    )

    wb = load_workbook(xlsx_path)
    try:
        ws = wb.active
        ws.cell(row=2, column=1).font = Font(name="Calibri", bold=True, color="00FF0000")
        wb.save(xlsx_path)
    finally:
        wb.close()

    update_row(xlsx_path, 0, "a_t_lp.jpg")

    wb = load_workbook(xlsx_path)
    try:
        ws = wb.active
        image_cell = ws.cell(row=2, column=3)
        assert image_cell.value == "a_t_lp.jpg"
        assert image_cell.font.bold is True
        assert image_cell.font.name == "Calibri"
    finally:
        wb.close()
