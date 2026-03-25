"""CSV reader / writer with in-place row update and resume support.

Design goals
------------
* Process rows **sequentially** – one at a time.
* **Resumable** – if the process is interrupted the CSV is already updated up
  to the last completed row, so re-running starts from where we left off.
* A row is considered *done* when the ``image_file`` column is non-empty **and**
  the referenced file exists on disk.  Otherwise the row is (re-)processed.
* All changes are written back to the **same** CSV file (in-place).
"""

from __future__ import annotations

import csv
from copy import copy
import logging
import os
import re
import tempfile
import time
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook

logger = logging.getLogger(__name__)

# The name of the column we add / update
IMAGE_COLUMN = "image_file"
EXCEL_IMAGE_COLUMN = "Image"

# Characters not allowed in file-system names
_UNSAFE_CHARS = re.compile(r'[\\/:*?"<>|]')
_CSV_DELIMITERS = ",;\t|"


# ---------------------------------------------------------------------------
# Public helpers
# ---------------------------------------------------------------------------


def sanitize_filename(text: str) -> str:
    """Replace unsafe filename characters with underscores and collapse spaces."""
    safe = _UNSAFE_CHARS.sub("_", text)
    # Collapse multiple whitespace / underscores
    safe = re.sub(r"[\s_]+", "_", safe).strip("_")
    return safe


def build_image_filename(artist: str, title: str, format_: str) -> str:
    """Return the standardised image filename for a record.

    Pattern: ``{artist}_{title}_{format}.jpg``
    All components are sanitised so the result is always a valid filename.
    """
    parts = [sanitize_filename(p) for p in [artist, title, format_] if p and p.strip()]
    return "_".join(parts) + ".jpg"


def reserve_unique_filename(
    filename: str,
    output_dir: str | Path,
    reserved_names: set[str],
    existing_names: set[str] | None = None,
) -> str:
    """Return a filename unique across output_dir and reserved_names.

    If ``filename`` already exists on disk (or in ``existing_names``) or was reserved for another pending
    row in the current run, numeric suffixes ``_2``, ``_3``, ... are appended.
    The chosen name is added to ``reserved_names`` before returning.
    """
    output_dir = Path(output_dir)
    path = Path(filename)
    stem = path.stem
    suffix = path.suffix or ".jpg"

    candidate = f"{stem}{suffix}"
    counter = 2
    while (
        candidate in reserved_names
        or (existing_names is not None and candidate in existing_names)
        or (existing_names is None and (output_dir / candidate).exists())
    ):
        candidate = f"{stem}_{counter}{suffix}"
        counter += 1

    reserved_names.add(candidate)
    if existing_names is not None:
        existing_names.add(candidate)
    return candidate


def iter_pending_rows(
    csv_path: str | Path,
    output_dir: str | Path,
    scan_log_every: int = 0,
) -> Iterator[tuple[int, dict[str, Any]]]:
    """Yield ``(row_index, row_dict)`` for every row that still needs processing.

    A row is considered done when:
    1. It has a non-empty ``image_file`` column, **and**
    2. The file ``output_dir / image_file`` exists on disk.

    Args:
        csv_path: Path to the input / output CSV file.
        output_dir: Folder where downloaded images are stored.
        scan_log_every: Emit scan progress log every N rows (0 disables).

    Yields:
        ``(row_index, row_dict)`` tuples (0-based row index, excluding header).
    """
    csv_path = Path(csv_path)
    output_dir = Path(output_dir)
    image_column = _get_image_column(csv_path)

    if _is_excel_path(csv_path):
        yield from _iter_pending_rows_excel(
            csv_path,
            output_dir,
            image_column,
            scan_log_every,
        )
        return

    rows = _read_all_rows(csv_path)
    for idx, row in enumerate(rows):
        if scan_log_every > 0 and (idx + 1) % scan_log_every == 0:
            logger.info("Scanning rows... %d checked", idx + 1)

        # Stop at the first fully-blank data row (end-of-data sentinel)
        if _is_blank_row(row, image_column):
            logger.info("Blank row encountered at index %d – stopping scan.", idx)
            break

        image_file = (row.get(image_column, "") or "").strip()
        if image_file and (output_dir / image_file).exists():
            logger.debug("Row %d already done (%s) – skipping", idx, image_file)
            continue
        yield idx, row


def update_row(
    csv_path: str | Path,
    row_index: int,
    image_filename: str,
) -> None:
    """Set ``image_file`` on the given row and save the CSV atomically.

    Args:
        csv_path: Path to the CSV file.
        row_index: 0-based data row index (not counting the header).
        image_filename: Value to write into the ``image_file`` column.
    """
    update_rows(csv_path, {row_index: image_filename})

    logger.debug("Updated row %d → %s", row_index, image_filename)


def update_rows(
    csv_path: str | Path,
    updates: dict[int, str],
) -> None:
    """Set image filename for multiple rows and save atomically.

    Args:
        csv_path: Path to CSV/XLSX file.
        updates: Mapping of 0-based row_index to image filename.
    """
    if not updates:
        return

    csv_path = Path(csv_path)
    sorted_updates = sorted(updates.items(), key=lambda item: item[0])

    if _is_excel_path(csv_path):
        logger.info(
            "Loading XLSX workbook for batch update (%d rows): %s",
            len(sorted_updates),
            csv_path,
        )
        _update_excel_rows(csv_path, sorted_updates)
        logger.info("Finished XLSX batch update: %s", csv_path)
        return

    rows = _read_all_rows(csv_path)
    for row_index, image_filename in sorted_updates:
        if row_index >= len(rows):
            raise IndexError(f"Row index {row_index} out of range ({len(rows)} rows)")
        rows[row_index][IMAGE_COLUMN] = image_filename

    _write_all_rows(csv_path, rows)


def ensure_image_column(csv_path: str | Path) -> None:
    """Add the ``image_file`` column to the CSV if it is not already present.

    This is a no-op if the column already exists.  It is safe to call on every
    program start before the main loop.
    """
    csv_path = Path(csv_path)

    if _is_excel_path(csv_path):
        headers = _read_excel_headers(csv_path)
        if EXCEL_IMAGE_COLUMN not in headers:
            raise ValueError(
                f"Excel file {csv_path} must contain an existing '{EXCEL_IMAGE_COLUMN}' column."
            )
        return

    rows = _read_all_rows(csv_path)
    if rows and IMAGE_COLUMN not in rows[0]:
        for row in rows:
            row.setdefault(IMAGE_COLUMN, "")
        _write_all_rows(csv_path, rows)
        logger.info("Added '%s' column to %s", IMAGE_COLUMN, csv_path)


# ---------------------------------------------------------------------------
# Private helpers
# ---------------------------------------------------------------------------


def _read_all_rows(csv_path: Path) -> list[dict[str, Any]]:
    if _is_excel_path(csv_path):
        return _read_all_rows_excel(csv_path)

    delimiter = _detect_delimiter(csv_path)
    with csv_path.open("r", newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh, delimiter=delimiter, quotechar='"', escapechar='\\')
        return [dict(row) for row in reader]


def _read_all_rows_excel(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        rows_iter = list(ws.iter_rows(values_only=True))
        if not rows_iter:
            return []

        headers = [_normalize_header(cell) for cell in rows_iter[0]]
        data_rows: list[dict[str, Any]] = []
        for row_values in rows_iter[1:]:
            row_dict = {
                header: _normalize_cell_value(value)
                for header, value in zip(headers, row_values)
                if header
            }
            if row_dict:
                data_rows.append(row_dict)

        return data_rows
    finally:
        wb.close()


def _iter_pending_rows_excel(
    path: Path,
    output_dir: Path,
    image_column: str,
    scan_log_every: int,
) -> Iterator[tuple[int, dict[str, Any]]]:
    """Stream XLSX rows to avoid full in-memory loading before processing."""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)

        header_values = next(rows_iter, None)
        if header_values is None:
            return

        headers = [_normalize_header(cell) for cell in header_values]

        for idx, row_values in enumerate(rows_iter):
            if scan_log_every > 0 and (idx + 1) % scan_log_every == 0:
                logger.info("Scanning spreadsheet rows... %d checked", idx + 1)

            row = {
                header: _normalize_cell_value(value)
                for header, value in zip(headers, row_values)
                if header
            }

            # Stop at the first fully-blank data row (end-of-data sentinel)
            if _is_blank_row(row, image_column):
                logger.info("Blank row encountered at index %d – stopping scan.", idx)
                break

            image_file = (row.get(image_column, "") or "").strip()
            if image_file and (output_dir / image_file).exists():
                continue
            yield idx, row
    finally:
        wb.close()


def _write_all_rows(csv_path: Path, rows: list[dict[str, Any]]) -> None:
    """Write *rows* back to *csv_path* atomically (temp-file + rename)."""
    if not rows:
        return

    delimiter = _detect_delimiter(csv_path)

    fieldnames = list(rows[0].keys())
    if IMAGE_COLUMN not in fieldnames:
        fieldnames.append(IMAGE_COLUMN)

    # Write to a sibling temp file then rename for atomicity
    dir_ = csv_path.parent
    fd, tmp_path = tempfile.mkstemp(dir=dir_, prefix=".tmp_", suffix=".csv")
    try:
        with os.fdopen(fd, "w", newline="", encoding="utf-8") as fh:
            writer = csv.DictWriter(
                fh,
                fieldnames=fieldnames,
                extrasaction="ignore",
                delimiter=delimiter,
                quotechar='"',
                escapechar='\\',
            )
            writer.writeheader()
            writer.writerows(rows)
        os.replace(tmp_path, csv_path)
    except Exception:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass
        raise


def _update_excel_row(path: Path, row_index: int, image_filename: str) -> None:
    _update_excel_rows(path, [(row_index, image_filename)])


def _update_excel_rows(path: Path, updates: list[tuple[int, str]]) -> None:
    wb = load_workbook(path)
    try:
        ws = wb.active
        headers = [_normalize_header(cell.value) for cell in ws[1]]
        try:
            image_col_idx = headers.index(EXCEL_IMAGE_COLUMN) + 1
        except ValueError as exc:
            raise ValueError(
                f"Excel file {path} must contain an existing '{EXCEL_IMAGE_COLUMN}' column."
            ) from exc

        for row_index, image_filename in updates:
            sheet_row = row_index + 2
            if sheet_row > ws.max_row:
                raise IndexError(f"Row index {row_index} out of range ({max(0, ws.max_row - 1)} rows)")

            cell = ws.cell(row=sheet_row, column=image_col_idx)
            original_style = _resolve_cell_style_template(ws, sheet_row, image_col_idx)
            cell.value = image_filename
            if original_style is not None:
                cell._style = original_style

        fd, tmp_path = tempfile.mkstemp(dir=path.parent, prefix=".tmp_", suffix=path.suffix)
        os.close(fd)
        os.unlink(tmp_path)
        wb.save(tmp_path)
        _replace_with_retries(Path(tmp_path), path)
    finally:
        wb.close()


def _replace_with_retries(tmp_path: Path, target_path: Path, retries: int = 3, delay: float = 0.5) -> None:
    """Replace a file with small retries to tolerate brief Windows file locks."""
    last_error: Exception | None = None
    for attempt in range(1, retries + 1):
        try:
            os.replace(tmp_path, target_path)
            return
        except PermissionError as exc:
            last_error = exc
            if attempt < retries:
                logger.warning(
                    "Spreadsheet file is locked, retrying replace (%d/%d): %s",
                    attempt,
                    retries,
                    target_path,
                )
                time.sleep(delay)
            else:
                break

    try:
        tmp_path.unlink(missing_ok=True)
    except OSError:
        pass

    raise PermissionError(
        f"Could not update spreadsheet {target_path}. It may be open in Excel or locked by another process."
    ) from last_error


def _resolve_cell_style_template(ws: Any, row: int, col: int) -> Any | None:
    """Return a style template for a target cell, preferring same-row styles.

    Excel files often use row/table styling where the target image cell itself has no
    explicit style assigned yet. In that case we inherit style from another styled
    cell in the same row so writing a value does not visually change formatting.
    """
    target = ws.cell(row=row, column=col)
    if target.has_style:
        return copy(target._style)

    for probe_col in range(1, ws.max_column + 1):
        probe = ws.cell(row=row, column=probe_col)
        if probe.has_style:
            return copy(probe._style)

    for probe_row in (row - 1, row + 1):
        if probe_row < 1 or probe_row > ws.max_row:
            continue
        probe = ws.cell(row=probe_row, column=col)
        if probe.has_style:
            return copy(probe._style)

    return None


def _read_excel_headers(path: Path) -> list[str]:
    logger.info("Loading XLSX headers: %s", path)
    wb = load_workbook(path, data_only=True)
    try:
        ws = wb.active
        headers = [_normalize_header(cell.value) for cell in ws[1] if _normalize_header(cell.value)]
        logger.info("Loaded %d XLSX headers from %s", len(headers), path)
        return headers
    finally:
        wb.close()


def _normalize_header(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _normalize_cell_value(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _is_blank_row(row: dict[str, Any], image_column: str) -> bool:
    """Return True if *row* contains no meaningful data (end-of-spreadsheet sentinel).

    A row is considered blank when every column value is empty, ignoring the
    image-filename column which may be legitimately empty before processing.
    """
    for key, value in row.items():
        if key == image_column:
            continue
        if (value or "").strip():
            return False
    return True


def _is_excel_path(path: Path) -> bool:
    return path.suffix.lower() in {".xlsx", ".xlsm", ".xltx", ".xltm"}


def _get_image_column(path: Path) -> str:
    return EXCEL_IMAGE_COLUMN if _is_excel_path(path) else IMAGE_COLUMN


def _detect_delimiter(csv_path: Path) -> str:
    """Detect delimiter from CSV content; fallback to comma."""
    with csv_path.open("r", newline="", encoding="utf-8-sig") as fh:
        sample = fh.read(4096)

    if not sample:
        return ","

    header = sample.splitlines()[0] if sample.splitlines() else ""
    header_delim_counts = {d: header.count(d) for d in _CSV_DELIMITERS}
    dominant_header_delim = max(header_delim_counts, key=header_delim_counts.get)
    if header_delim_counts[dominant_header_delim] > 0:
        return dominant_header_delim

    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=_CSV_DELIMITERS)
        return dialect.delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in _CSV_DELIMITERS}
        delimiter = max(counts, key=counts.get)
        return delimiter if counts[delimiter] > 0 else ","
